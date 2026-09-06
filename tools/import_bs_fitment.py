#!/usr/bin/env python3
"""Import complete rows from the user-designated official BS fitment workbook."""
import argparse, hashlib, json, re, unicodedata
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
MAKERS = {'レクサス':'レクサス','トヨタ':'トヨタ','日産':'日産','ホンダ':'ホンダ','マツダ':'マツダ',
          'スバル':'SUBARU','SUBARU':'SUBARU','スズキ':'スズキ','ダイハツ':'ダイハツ','三菱':'三菱'}

def norm(v):
    s=unicodedata.normalize('NFKC',str(v or '')).lower()
    s=''.join(chr(ord(c)-0x60) if 'ァ'<=c<='ヶ' else c for c in s)
    return re.sub(r'[\s\-‐‑‒–—―・_/ー：:（）()．.]','',s)

def period(v):
    s=unicodedata.normalize('NFKC',str(v or '')).replace('〜','~').replace('～','~')
    m=re.search(r'(\d{2,4})/(\d{1,2})\s*~\s*(?:(\d{2,4})/(\d{1,2}))?',s)
    if not m:return None
    def year(x):
        y=int(x);return 2000+y if y<100 else y
    return f'{year(m[1]):04d}-{int(m[2]):02d}', (f'{year(m[3]):04d}-{int(m[4]):02d}' if m[3] else '2025-12')

def main():
    ap=argparse.ArgumentParser();ap.add_argument('workbook');ap.add_argument('--linked',required=True);args=ap.parse_args()
    path=Path(args.workbook);sha=hashlib.sha256(path.read_bytes()).hexdigest()
    master=json.loads((ROOT/'app/data/jp_vehicle_search_master_2000_2026_v1.json').read_text())['vehicles']
    linked=set(json.loads(Path(args.linked).read_text()))
    wb=load_workbook(path,read_only=True,data_only=True);ws=wb['装着参考例']
    current_maker=None;out=[];evidence=[]
    for rn,row in enumerate(ws.iter_rows(min_row=4,values_only=True),4):
        vals=list(row);first=str(vals[0] or '').strip()
        if first in MAKERS and not vals[1]:current_maker=MAKERS[first];continue
        if not current_maker or not (first and vals[1] and vals[2]):continue
        name=first;per=period(vals[1]);code=unicodedata.normalize('NFKC',str(vals[2])).strip()
        tire=unicodedata.normalize('NFKC',str(vals[4] or '')).upper().replace('RF','R')
        tire=re.sub(r'^[FR][：:]','',tire).strip()
        pcdm=re.fullmatch(r'([456])-\s*(100|108|110|114(?:\.3)?|120|127|139(?:\.7)?|150)',unicodedata.normalize('NFKC',str(vals[6] or '')).strip())
        hubm=re.fullmatch(r'\d+(?:\.\d+)?',unicodedata.normalize('NFKC',str(vals[7] or '')).strip())
        fast=re.search(r'M?(12|14)\s*[X×x]\s*(1\.25|1\.5)',unicodedata.normalize('NFKC',str(vals[8] or '')),re.I)
        tm=re.fullmatch(r'(\d{3}/\d{2}R\d{2})',tire)
        if not (per and pcdm and hubm and fast and tm):continue
        source_name=norm(name);matches=[]
        for m in master:
            if m['search_id'] in linked or m['maker']!=current_maker:continue
            aliases={norm(m['model']),*(norm(a) for a in m.get('aliases',[]))}
            score=max((len(a) for a in aliases if len(a)>=2 and (source_name.startswith(a) or a.startswith(source_name))),default=0)
            if score:matches.append((score,m))
        if matches:
            best=max(x[0] for x in matches);matches=[x[1] for x in matches if x[0]==best]
        if len(matches)!=1:continue
        m=matches[0];vid='BS_'+hashlib.sha256(f'{sha}:{rn}:{m["search_id"]}'.encode()).hexdigest()[:12].upper()
        raw=[x for x in vals[:11]]
        ev=dict(id=vid,source_file=path.name,sha256=sha,sheet=ws.title,row=rn,raw_columns=raw,checked_at='2026-09-06')
        evidence.append(ev)
        method='bolt' if 'ボルト' in str(vals[9] or '') else 'nut'
        out.append(dict(vehicle_id=vid,maker=m['maker'],model=m['model'],generation=code,
            model_codes=[x for x in re.findall(r'[A-Z]{1,6}\d{1,4}[A-Z]{0,3}',code)][:12] or [code],
            year_from=per[0],year_to=per[1],pcd=float(pcdm[2]),holes=int(pcdm[1]),hub_bore=float(hubm[0]),
            fastener=f'M{fast[1]}×P{fast[2]}',fastener_details=dict(method=method,thread_diameter='M'+fast[1],thread_pitch=float(fast[2])),
            oem_inch=tm[1][-2:],oem_tire=tm[1],confidence='C',
            notes='ユーザー指定の公式BS装着参考例を転記。終了月なしの年式は資料収録期限2025-12までを検索範囲として登録。装着前に現車・グレード・荷重・キャリパーを確認。'+((' 原文備考: '+str(vals[9])) if vals[9] else ''),
            sources=[dict(source_type='user_designated_official_matching_document',source_name=f'{path.name} {ws.title} 行{rn}',
              source_url='https://tire-wheel-price-navi.hide718283.chatgpt.site/data/vehicle-updates/bs-source-excerpts-2026-09-06.json#'+vid,
              verified_at='2026-09-06',document_sha256=sha,document_sheet=ws.title,document_row=rn)],source_document=ev))
    dest=ROOT/'app/data/vehicle-updates/supplied-bs-2026-09-06.json'
    dest.write_text(json.dumps(dict(schema_version='1.0.0',updates=out),ensure_ascii=False,indent=2)+'\n')
    (ROOT/'app/data/vehicle-updates/bs-source-excerpts-2026-09-06.json').write_text(json.dumps(dict(source_file=path.name,sha256=sha,rows=evidence),ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(dict(records=len(out),models=len({(x['maker'],x['model']) for x in out})),ensure_ascii=False))
    print(' / '.join(x['model'] for x in out))

if __name__=='__main__':main()
