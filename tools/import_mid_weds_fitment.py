#!/usr/bin/env python3
"""Join user-designated official MID rows with exact Weds model fitment pages."""
import argparse, hashlib, json, re, unicodedata
from pathlib import Path
from openpyxl import load_workbook
from research_all_fitment import ROOT, norm

SHEETS={'レクサス':'レクサス','トヨタ':'トヨタ','ニッサン':'日産','ホンダ':'ホンダ','マツダ':'マツダ','ミツビシ':'三菱','スバル':'SUBARU','スズキ':'スズキ','ダイハツ':'ダイハツ','４ＷＤ':'*'}

def dates(v):
 s=unicodedata.normalize('NFKC',str(v or '')).replace('〜','~').replace('～','~')
 m=re.search(r'(20\d{2})/(\d{1,2})\s*~\s*(?:(20\d{2})/(\d{1,2}))?',s)
 return (f'{m[1]}-{int(m[2]):02d}',f'{m[3]}-{int(m[4]):02d}' if m[3] else '2025-12') if m else None

def main():
 ap=argparse.ArgumentParser();ap.add_argument('workbook');ap.add_argument('--linked',required=True);ap.add_argument('--weds',required=True);args=ap.parse_args()
 p=Path(args.workbook);sha=hashlib.sha256(p.read_bytes()).hexdigest();linked=set(json.load(open(args.linked)))
 master=json.loads((ROOT/'app/data/jp_vehicle_search_master_2000_2026_v1.json').read_text())['vehicles']
 weds=json.load(open(args.weds))['weds'];wb=load_workbook(p,read_only=True,data_only=True);out=[];ev=[]
 for title,maker in SHEETS.items():
  ws=wb[title]
  for rn,row in enumerate(ws.iter_rows(min_row=3,values_only=True),3):
   v=list(row);name=str(v[0] or '').strip();ds=dates(v[2] if len(v)>2 else None)
   pm=re.fullmatch(r'([456])\s*[/／]\s*(100|108|110|112|114(?:\.3)?|120|127|130|139(?:\.7)?|150)',unicodedata.normalize('NFKC',str(v[3] or '')).strip()) if len(v)>3 else None
   if not(name and v[1] and ds and pm):continue
   inch=str(v[6] or '').strip();ts=unicodedata.normalize('NFKC',str(v[7] or '')).upper().strip()
   tire_part=re.search(r'\d{3}/\d{2}',ts)
   if not re.fullmatch(r'\d{2}',inch) or not tire_part:continue
   tire=tire_part[0]+'R'+inch;source_parts={norm(name),*(norm(x) for x in re.split(r'[/／]',name) if x)}
   for m in master:
    if m['search_id'] in linked or (maker!='*' and m['maker']!=maker and not (m['maker']=='トヨタ' and m['model']=='コペン' and maker=='ダイハツ')):continue
    aliases={norm(m['model']),*(norm(a) for a in m.get('aliases',[]))}
    if not (source_parts & aliases) and not any(len(a)>=3 and a in norm(name) for a in aliases):continue
    target_weds=[x for x in weds if x['maker']==m['maker'] and ({norm(x['model']),*(norm(a) for a in x['model'].split('/'))}&aliases)
      and x.get('holes')==int(pm[1]) and x.get('pcd') in (float(pm[2]),{114:114.3,139:139.7}.get(float(pm[2])))
      and x.get('hub_bore') and x.get('thread_diameter') and x.get('thread_pitch') and x.get('method')]
    specs={(x['pcd'],x['holes'],x['hub_bore'],x['thread_diameter'],x['thread_pitch'],x['method']) for x in target_weds}
    if len(specs)!=1:continue
    pcd,holes,hub,diam,pitch,method=next(iter(specs));wid=target_weds[0]
    vid='MID_'+hashlib.sha256(f'{sha}:{title}:{rn}:{m["search_id"]}'.encode()).hexdigest()[:12].upper()
    evidence=dict(id=vid,source_file=p.name,sha256=sha,sheet=title,row=rn,raw_columns=v[:9],checked_at='2026-09-06');ev.append(evidence)
    out.append(dict(vehicle_id=vid,maker=m['maker'],model=m['model'],generation=unicodedata.normalize('NFKC',str(v[1])).strip(),model_codes=[unicodedata.normalize('NFKC',str(v[1])).strip()],
      year_from=ds[0],year_to=ds[1],pcd=pcd,holes=holes,hub_bore=hub,fastener=f'M{diam}×P{pitch:g}',
      fastener_details=dict(method='bolt' if method=='ボルト' else 'nut',thread_diameter=f'M{diam}',thread_pitch=pitch),oem_inch=inch,oem_tire=tire,confidence='B',
      notes='ユーザー指定の公式MID適合表の型式・年式・PCD・純正サイズを、同一車名・同一PCDのWeds公式適合ページのハブ径・締結規格と照合。装着前に現車・グレードを確認。',
      sources=[dict(source_type='user_designated_official_matching_document',source_name=f'{p.name} {title} 行{rn}',source_url='https://tire-wheel-price-navi.hide718283.chatgpt.site/data/vehicle-updates/mid-source-excerpts-2026-09-06.json#'+vid,verified_at='2026-09-06',document_sha256=sha,document_sheet=title,document_row=rn),
               dict(source_type='wheel_manufacturer_official',source_name='Weds '+wid['model']+' '+wid['generation'],source_url=wid.get('source_url',wid['url']),verified_at='2026-09-06')],source_document=evidence))
 dest=ROOT/'app/data/vehicle-updates/supplied-mid-weds-2026-09-06.json';dest.write_text(json.dumps(dict(schema_version='1.0.0',updates=out),ensure_ascii=False,indent=2)+'\n')
 (ROOT/'app/data/vehicle-updates/mid-source-excerpts-2026-09-06.json').write_text(json.dumps(dict(source_file=p.name,sha256=sha,rows=ev),ensure_ascii=False,indent=2)+'\n')
 print(json.dumps(dict(records=len(out),models=len({x['model'] for x in out})),ensure_ascii=False));print(' / '.join(x['model'] for x in out))
if __name__=='__main__':main()
